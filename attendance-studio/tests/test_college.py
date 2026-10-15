
import io

import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pdfplumber
from reportlab.pdfgen import canvas

from backend.college import activities, cdc_pdf, cdc_input, compare_college, odd_students
from backend.main import app



def cdc_fixture():
    output = io.BytesIO(); pdf = canvas.Canvas(output)
    for y, text in [(770,'FN Attendance Summary - 29-09-2026'), (745,'B.Tech III Semester Attendance Summary'), (720,'BATCH BRANCH Total Strength Present Absent'), (695,'SU1 CSE 3 2 1')]:
        pdf.drawString(40,y,text)
    pdf.showPage()
    for y, text in [(770,'III SEM CSE - Absentees List'), (745,'S.No Roll No Name Branch Batch'), (720,'1 25951A0501 TEST STUDENT CSE SU1'), (695,'Total Absent: 1')]:
        pdf.drawString(40,y,text)
    pdf.save();return output.getvalue()


def class_fixture():
    return b'Date,Period,Section,Absentees,Staff\n2026-09-29,1,CSE-III-SEM-A,25951A0502,IARE1000 - Dr. Test Faculty\n2026-09-29,2,CSE-III-SEM-A,NIL,IARE1000 - Dr. Test Faculty\n'


def index_fixture():
    return {'warnings': [], 'students': {f'25951A050{n}': {'roll':f'25951A050{n}', 'name':'Student '+str(n), 'branch':'CSE','batch':'SU1','semester':3,'section':'A'} for n in (1,2,3)}}


class CollegeTests(unittest.TestCase):
    def test_only_cdc_present_class_absent_students_are_discrepancies(self):
        index = index_fixture()
        index['students']['25951A0503']['name'] = ''
        index['students']['25951A0504'] = dict(index['students']['25951A0503'], roll='25951A0504', batch='')
        cdc = cdc_pdf(cdc_fixture(), 'cdc.pdf')
        rows = activities(class_fixture(), 'class.csv')[:1]
        rows[0]['absentees'] = ['25951A0501', '25951A0503', '25951A0503', '25951A0504', '25951A0599']
        with patch('backend.college.load_index', return_value=index):
            report = compare_college(cdc, rows, '2026-09-29')
        self.assertEqual(report['rows'][0]['department_absent'], 2)
        self.assertEqual(report['rows'][0]['discrepancy'], 1)
        self.assertEqual([r['Roll Number'] for r in report['details'] if r['Discrepancy']=='Yes'], ['25951A0503'])
        rows[0]['absentees'] = ['25951A0501']
        with patch('backend.college.load_index', return_value=index):
            report = compare_college(cdc, rows, '2026-09-29')
        self.assertEqual(report['rows'][0]['discrepancy'], 0)

    def test_non_cdc_students_never_enter_comparison_or_odd_list(self):
        student_index=index_fixture()
        student_index['students']['25951A0504']={'roll':'25951A0504','name':'Non CDC','branch':'CSE','batch':'Unspecified','semester':3,'section':'A'}
        student_index['students']['25951A0503']['section']=''
        student_index['warnings']=['25951A0502: conflicting master batch/branch.']
        with patch('backend.college.load_index',return_value=student_index):
            odd=odd_students(student_index)
            self.assertEqual({r['roll'] for r in odd},{'25951A0502','25951A0503'})
            student_index['warnings'] = []
            register=b'Date,Session,Semester,Roll Number,Status\n2026-09-29,FN,III,25951A0501,Present\n2026-09-29,FN,III,25951A0502,Absent\n2026-09-29,FN,III,25951A0504,Absent\n'
            cdc=cdc_input(register,'cdc.csv')
            self.assertNotIn('25951A0504',cdc['cohort_rolls'])
            self.assertNotIn('25951A0504',cdc['absentees'])
            class_records=activities(b'Date,Period,Section,Absentees,Staff\n2026-09-29,1,CSE-III-SEM-A,"25951A0502,25951A0504",IARE1000 - Dr. Test Faculty\n','class.csv')
            result=compare_college(cdc,class_records,'2026-09-29')
            self.assertEqual(result['rows'][0]['department_absent'],1)
            self.assertNotIn('25951A0504',{r['Roll Number'] for r in result['details']})
            partial = compare_college(cdc | {'absentees': {'25951A0504': {'batch': 'SU1', 'branch': 'CSE'}}},class_records,'2026-09-29')
            self.assertIn('25951A0504', {r['Roll Number'] for r in partial['excluded']})
            activity=b'Date,Period,Section,Absentees,Staff\n2026-09-29,1,CSE-III-SEM-A,25951A0504,IARE1000 - Dr. Test Faculty\n2026-09-29,2,CSE-III-SEM-A,NIL,IARE1000 - Dr. Test Faculty\n'
            self.assertEqual(cdc_input(activity,'cdc.csv')['absentees'],{})

    def test_cdc_excel_and_class_pdf(self):
        from openpyxl import Workbook
        from reportlab.platypus import Table, TableStyle
        book=Workbook(); sheet=book.active
        sheet.append(['Date','Session','Semester','Roll Number','Status'])
        for n in (1,2,3):
            sheet.append(['2026-09-29','FN','III',f' 25951a050{n} ', 'Absent' if n==1 else 'Present'])
        sheet.append(['2026-09-29','FN','III','25951A0501','Absent'])
        excel=io.BytesIO();book.save(excel)
        output=io.BytesIO();pdf=canvas.Canvas(output,pagesize=(900,600))
        for hour in (1,2):
            table=Table([['Date','Period','Section','Absentees','Staff'],['2026-09-29',str(hour),'CSE-III-SEM-A','25951A0502','IARE1000 - Dr. Test Faculty']],colWidths=[110,60,150,160,250])
            table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,(0,0,0))]))
            table.wrapOn(pdf,850,500);table.drawOn(pdf,40,450);pdf.showPage()
        pdf.save()
        with patch('backend.college.load_index', return_value=index_fixture()):
            left=cdc_input(excel.getvalue(),'cdc.xlsx')
            self.assertEqual(len(left['cohort_rolls']),3)
            self.assertEqual(len(left['absentees']),1)
            right=activities(output.getvalue(),'class.pdf')
            self.assertEqual(len(right),2)
            result=compare_college(left,right,'2026-09-29')
            self.assertEqual(result['rows'][0]['class_only'],1)
            self.assertEqual(result['rows'][0]['cdc_only'],1)
            self.assertEqual(len(result['details']),4)
            client=TestClient(app)
            response=client.post('/api/college/compare',files=[('cdc',('cdc.xlsx',excel.getvalue())),('classroom',('class.pdf',output.getvalue()))],data={'report_date':'2026-09-29'})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(client.post('/api/college/inspect-cdc',files={'file':('cdc.xlsx',excel.getvalue())}).status_code,200)
            self.assertEqual(client.post('/api/college/inspect-class',files={'file':('class.pdf',output.getvalue())}).status_code,200)

    def test_cdc_activity_and_invalid_pdf(self):
        with patch('backend.college.load_index', return_value=index_fixture()):
            single=class_fixture().splitlines()[:2]
            result=cdc_input(b'\n'.join(single),'cdc.csv')
            self.assertEqual(set(result['absentees']),{'25951A0502'})
            with self.assertRaisesRegex(ValueError,'different CDC student absentee'):
                cdc_input(class_fixture(),'cdc.csv')
        with self.assertRaisesRegex(ValueError,'no readable activity table'):
            activities(cdc_fixture(),'class.pdf')

    def test_equal_counts_still_find_different_students(self):
        with patch('backend.college.load_index', return_value=index_fixture()):
            result = compare_college(cdc_pdf(cdc_fixture(),'cdc.pdf'),activities(class_fixture(),'class.csv'),'2026-09-29')
        row = result['rows'][0]
        self.assertEqual((row['discrepancy'],row['cdc_only'],row['class_only']), (1,1,1))
        self.assertEqual([r['Missing Attendance'] for r in result['details'][:3]], ['Class','None','None'])
        self.assertEqual(len(result['details']),4)

    def test_date_semester_and_database_guards(self):
        cdc = cdc_pdf(cdc_fixture(),'cdc.pdf'); rows = activities(class_fixture(),'class.csv')
        with self.assertRaisesRegex(ValueError,'dated'):
            compare_college(cdc,rows,'2026-09-30')
        with self.assertRaisesRegex(ValueError,'same date and semester'):
            compare_college(cdc,[r | {'semester':5} for r in rows],'2026-09-29')
        with patch('backend.college.load_index', return_value=index_fixture()):
            result = compare_college(cdc,[rows[0] | {'absentees':['25951A0599']}],'2026-09-29')
            self.assertFalse(result['partial'])
            self.assertEqual(result['rows'][0]['department_absent'], 0)
            self.assertNotIn('25951A0599', {r['Roll Number'] for r in result['details']})
            with self.assertRaisesRegex(ValueError,'Conflicting'):
                compare_college(cdc,[rows[0],rows[0] | {'faculty_id':'IARE2000'}],'2026-09-29')

    def test_discrepancy_is_absolute_in_exports(self):
        from backend.college_reports import excel_report, pdf_report
        cdc = cdc_pdf(cdc_fixture(), 'cdc.pdf')
        cdc['absentees'] = {}
        with patch('backend.college.load_index', return_value=index_fixture()):
            result = compare_college(cdc, activities(class_fixture(), 'class.csv'), '2026-09-29')
        self.assertEqual(result['rows'][0]['discrepancy'], 1)
        self.assertEqual(result['rows'][0]['class_only'], 1)
        book = load_workbook(io.BytesIO(excel_report(result)))
        self.assertEqual(book['College Summary']['H11'].value, 1)
        self.assertEqual(book['College Summary']['H12'].value, 1)
        with pdfplumber.open(io.BytesIO(pdf_report(result))) as doc:
            self.assertNotIn('-1', '\n'.join(p.extract_text() for p in doc.pages))

    def test_downloads_and_filtered_totals(self):
        client = TestClient(app)
        with patch('backend.college.load_index', return_value=index_fixture()):
            response = client.post('/api/college/compare',files=[('cdc',('cdc.pdf',cdc_fixture(),'application/pdf')),('classroom',('class.csv',class_fixture(),'text/csv'))],data={'report_date':'2026-09-29'})
        self.assertEqual(response.status_code,200,response.text)
        payload = {'report_id':response.json()['report_id'],'branch':'CSE','hour':1}
        exported = client.post('/api/college/export?format=xlsx',json=payload)
        self.assertEqual(exported.status_code,200)
        book = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(book['College Summary']['H11'].value,1)
        self.assertEqual(book['College Summary']['F12'].value,1)
        self.assertEqual(book['Student Attendance'].max_row,3)
        self.assertEqual(book['Student Attendance']['C2'].value,'25951A0502')
        self.assertEqual(len(book['College Summary']._images),1)
        pdf = client.post('/api/college/export?format=pdf',json=payload)
        self.assertEqual(pdf.status_code,200)
        with pdfplumber.open(io.BytesIO(pdf.content)) as document:
            text = document.pages[0].extract_text()
            self.assertIn('29 September, 2026',text)
            self.assertIn('DEAN',text)
            self.assertIn('Dr. Test Faculty',text)
        self.assertEqual(client.post('/api/college/export',json={'report_id':'expired'}).status_code,422)

    def test_reported_unresolved_ids_are_excluded_and_exported(self):
        from backend.college_reports import excel_report, pdf_report
        index = index_fixture()
        missing = ['24951A0203', '24951A0204', '24951A0206', '24951A0212', '24951A0226', '24951A0231']
        index['students']['25955A0509'] = dict(index['students']['25951A0501'], roll='25955A0509', section='')
        index['students']['24951A0201'] = dict(index['students']['25951A0501'], roll='24951A0201', branch='EEE')
        cdc = cdc_pdf(cdc_fixture(), 'cdc.pdf')
        cdc['absentees']['25955A0509'] = {'batch':'SU1', 'branch':'CSE'}
        cdc['coverage'].append({'batch':'SU1', 'branch':'EEE', 'total':1, 'absent':0})
        rows = activities(class_fixture(), 'class.csv')
        rows.append(dict(rows[0], branch='EEE', absentees=missing))
        with patch('backend.college.load_index', return_value=index):
            result = compare_college(cdc, rows, '2026-09-29')
        omitted = {'25955A0509'}
        self.assertEqual({r['Roll Number'] for r in result['excluded']}, omitted)
        self.assertFalse(omitted & {r['Roll Number'] for r in result['details']})
        self.assertEqual(result['rows'][-1]['department_absent'], 0)
        self.assertFalse(set(missing) & {r['Roll Number'] for r in result['details']})
        book = load_workbook(io.BytesIO(excel_report(result)))
        self.assertEqual(book['Excluded Records'].max_row, 2)
        with pdfplumber.open(io.BytesIO(pdf_report(result))) as document:
            text = '\n'.join(p.extract_text() for p in document.pages)
        self.assertNotIn('Partial report', text)
        self.assertNotIn('Report notes', text)
        self.assertNotIn('25955A0509', text)
        self.assertNotIn('24951A0203', str([r['Roll Number'] for r in result['details']]))
        for student in index['students'].values():
            student['section'] = ''
        with patch('backend.college.load_index', return_value=index):
            with self.assertRaisesRegex(ValueError, 'No valid CDC students'):
                compare_college(cdc, rows, '2026-09-29')



if __name__ == '__main__':
    unittest.main()
