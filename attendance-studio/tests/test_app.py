import io
import json
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.engine import records, reconcile, summary, excel, inspect
from backend.main import app, local_filter

ROOT = Path(__file__).resolve().parents[1]


class AttendanceTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.cdc = (ROOT / 'tests/fixtures/cdc-register.csv').read_bytes()
        self.classroom = (ROOT / 'tests/fixtures/class-register.csv').read_bytes()

    def compare(self, cdc=None, classroom=None, options=None):
        return self.client.post('/api/reconcile', files={'cdc': ('cdc.csv', cdc if cdc is not None else self.cdc), 'classroom': ('class.csv', classroom if classroom is not None else self.classroom)}, data={'options': json.dumps(options or {'same_scope': True})})

    def test_complete_registers_and_api(self):
        response = self.compare()
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual(data['summary'], {'total': 12, 'None': 5, 'CDC': 2, 'Class': 3, 'Both': 2})
        self.assertEqual(data['files'][0]['duplicates'], 1)
        filtered = self.client.post('/api/query', json={'rows': data['rows'], 'question': 'Show BT26 IT students who missed class'})
        self.assertEqual(filtered.status_code, 200, filtered.text)
        self.assertIn('1 matching students', filtered.json()['answer'])
        self.assertEqual(filtered.json()['filters']['missing'], 'Class')
        report = self.client.post('/api/export', json={'rows': data['rows'], 'notes': data['notes']})
        self.assertEqual(report.status_code, 200)
        book = load_workbook(io.BytesIO(report.content))
        self.assertEqual(book['Attendance'].max_row, 13)
        self.assertEqual(book['Attendance']['A2'].fill.fgColor.rgb, '00DBEAFE')
        self.assertEqual(book['Attendance'].freeze_panes, 'D2')

    def test_absentee_lists(self):
        options = {'same_scope': True, 'cdc': {'mode': 'absentees'}, 'classroom': {'mode': 'absentees'}}
        response = self.compare(b'Roll No\n 001 \n002\n002\n', b'Roll Number\n002\n003\n', options)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result['summary'], {'total': 3, 'None': 0, 'CDC': 1, 'Class': 1, 'Both': 1})
        self.assertTrue(result['incomplete_cohort'])
        self.assertEqual(result['rows'][0]['Roll Number'], '001')

    def test_empty_absentee_list(self):
        options = {'same_scope': True, 'cdc': {'mode': 'absentees'}, 'classroom': {'mode': 'absentees'}}
        self.assertEqual(self.compare(b'Roll No\n', b'Roll No\n001\n', options).json()['summary']['Class'], 1)

    def test_validation(self):
        for left, right, text in [
            (b'Roll No,Status\n001,P\n001,A\n', b'Roll No,Status\n001,P\n', 'Conflicting'),
            (b'Roll No,Status\n001,Late\n', b'Roll No,Status\n001,P\n', 'unrecognized'),
            (b'Roll No,Status\n001,P\n', b'Roll No,Status\n002,P\n', 'missing students'),
            (b'Roll No,Batch,Status\n001,BT25,P\n', b'Roll No,Batch,Status\n001,BT26,P\n', 'conflicting batch'),
            (b'Roll No,Status\n,P\n', b'Roll No,Status\n001,P\n', 'missing roll'),
        ]:
            with self.subTest(text=text):
                response = self.compare(left, right)
                self.assertEqual(response.status_code, 422)
                self.assertIn(text, response.json()['detail'])
        self.assertEqual(self.compare(options={'same_scope': False}).status_code, 422)
        self.assertEqual(self.compare(b'Roll No\n001\n').status_code, 422)

    def test_formats_sheets_and_header_offset(self):
        for name in ['cdc-register.xlsx', 'cdc-register.xls']:
            data = (ROOT / 'tests/fixtures' / name).read_bytes()
            result, meta = records(data, name, {})
            self.assertEqual(len(result), 12)
            self.assertEqual(result['26IT002']['status'], 'Present')
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            pd.DataFrame([['Report'], ['Roll No', 'Status'], [' ab001 ', 'P']]).to_excel(writer, sheet_name='Attendance', index=False, header=False)
            pd.DataFrame([['Roll No','Status'],['002','A']]).to_excel(writer,sheet_name='Other',index=False,header=False)
        info, _ = inspect(buffer.getvalue(), 'book.xlsx', 'Attendance')
        self.assertEqual(info['header_row'], 2)
        result, _ = records(buffer.getvalue(), 'book.xlsx', {'sheet': 'Attendance'})
        self.assertEqual(result['AB001']['status'], 'Present')

    def test_boolean_absent_and_inference(self):
        parsed, _ = records(b'Roll No,Absent\n26it001,1\n26IT002,0\n', 'a.csv', {'infer_ids': True})
        self.assertEqual(parsed['26IT001'], {'batch': 'BT26', 'branch': 'IT', 'status': 'Absent'})
        self.assertEqual(parsed['26IT002']['status'], 'Present')

    def test_formula_injection_and_empty_export(self):
        rows = reconcile({'=1+1': {'batch': '=NOW()', 'branch': '+cmd', 'status': 'Absent'}}, {}, 'absentees', 'absentees')
        book = load_workbook(io.BytesIO(excel(rows, ['=1+1'])))
        self.assertEqual(book['Attendance']['A2'].data_type, 's')
        self.assertEqual(book['Attendance']['C2'].data_type, 's')
        self.assertEqual(load_workbook(io.BytesIO(excel([], [])))['Attendance'].max_row, 1)

    def test_query_guardrails(self):
        rows = self.compare().json()['rows']
        for query in ['except IT', 'Show BT26 or BT25', 'write a poem', 'attendance rate over the last month', 'Show BT99 IT', 'Show BT26 absences last week', 'missed CDC and Class']:
            with self.assertRaises(ValueError):
                local_filter(query, rows)
        with patch.dict('os.environ', {'OPENAI_API_KEY':'', 'OPENAI_MODEL':''}):
            response = self.client.post('/api/query', json={'rows': rows, 'question':'summary', 'use_ai': True})
            self.assertEqual(response.status_code, 422)
        response = self.client.post('/api/query', json={'rows': rows, 'question':'discrepancies by branch'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(sum(g['total'] for g in response.json()['groups']), 5)

    def test_malformed_and_ambiguous(self):
        self.assertEqual(self.client.post('/api/inspect',files={'file':('bad.txt',b'test')}).status_code,422)
        self.assertEqual(self.client.post('/api/inspect',files={'file':('bad.xlsx',b'not excel')}).status_code,422)
        data=b'Roll No,CDC Attendance,Class Attendance\n001,P,A\n'
        info,_=inspect(data,'a.csv')
        self.assertIn('status', info['ambiguous'])

    def test_ai_filter_is_validated_and_counts_are_local(self):
        rows = self.compare().json()['rows']
        plan = {'batch':'BT26','branch':'IT','missing':'Class','group_by':'none','supported':True}
        response = httpx.Response(200, request=httpx.Request('POST','https://api.openai.com/v1/responses'), json={'output':[{'content':[{'type':'output_text','text':json.dumps(plan)}]}]})
        with patch.dict('os.environ',{'OPENAI_API_KEY':'test-key','OPENAI_MODEL':'test-model'}), patch('backend.main.httpx.post',return_value=response) as call:
            result = self.client.post('/api/query',json={'rows':rows,'question':'Find IT class-only absences in BT26','use_ai':True})
            self.assertEqual(result.status_code,200,result.text)
            self.assertIn('1 matching students',result.json()['answer'])
            sent = call.call_args.kwargs['json']
            self.assertNotIn('26IT002',json.dumps(sent))
            self.assertFalse(sent['store'])


if __name__ == '__main__':
    unittest.main()
