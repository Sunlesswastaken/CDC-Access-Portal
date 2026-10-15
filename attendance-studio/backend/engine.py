"""Deterministic parsing and reconciliation. No model calls in this module."""
import csv
import io
import re
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = ['Batch Code', 'Branch', 'Roll Number', 'CDC Attendance', 'Class Attendance', 'Missing Attendance']
ALIASES = {
    'roll': {'roll', 'rollno', 'rollnumber', 'studentid', 'studentrollno', 'regno', 'registrationnumber', 'register number'.replace(' ', ''), 'enrollmentno', 'enrolmentnumber', 'admissionno'},
    'batch': {'batch', 'batchcode', 'cohort', 'yearofadmission'},
    'branch': {'branch', 'department', 'dept', 'stream', 'programme', 'program'},
    'status': {'attendance', 'status', 'attendancestatus', 'cdcattendance', 'classattendance', 'present', 'absent', 'presence'},
}
PRESENT = {'p', 'present', 'attended', 'yes', 'y', '1', 'true'}
ABSENT = {'a', 'absent', 'abs', 'not present', 'not attended', 'no', 'n', '0', 'false'}
LIMIT = 20000


def key(value):
    return re.sub(r'[^a-z0-9]', '', str(value).lower())


def clean(value):
    if value is None or pd.isna(value):
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().upper()


def tables(data, filename):
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.csv', '.xlsx', '.xls', '.pdf'}:
        raise ValueError('Use an XLSX, XLS, CSV or text-based PDF file.')
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('Each file must be 10 MB or smaller.')
    try:
        if suffix == '.pdf':
            from .pdf_import import pdf_table
            return pdf_table(data, ALIASES, key, LIMIT)
        if suffix == '.csv':
            text = data.decode('utf-8-sig')
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=',;\t').delimiter
            except csv.Error:
                delimiter = ','
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
            if len(rows) > LIMIT + 21:
                raise ValueError(f'Maximum {LIMIT:,} student rows per file.')
            return {'CSV': pd.DataFrame(rows).fillna('')}
        if suffix == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                    raise ValueError('This workbook expands beyond the 80 MB processing limit.')
        book = pd.ExcelFile(io.BytesIO(data), engine='openpyxl' if suffix == '.xlsx' else 'xlrd')
        if len(book.sheet_names) > 20:
            raise ValueError('Use a workbook with at most 20 sheets.')
        return {name: book.parse(name, header=None, dtype=object, keep_default_na=False, nrows=LIMIT + 22).fillna('') for name in book.sheet_names}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Could not read the file. Check its format, encoding (UTF-8 for CSV), and password protection. PDFs must contain selectable text.') from exc


def inspect(data, filename, sheet=None):
    frames = tables(data, filename)
    candidates = []
    for name, frame in frames.items():
        if sheet and name != sheet:
            continue
        for index, row in frame.head(20).iterrows():
            keys = [key(v) for v in row]
            if any(v in ALIASES['roll'] for v in keys):
                score = sum(any(v in aliases for v in keys) for aliases in ALIASES.values())
                candidates.append((score, name, index, row))
    if not candidates:
        raise ValueError('No roll-number header found in the first 20 rows. Rename the ID column to “Roll Number”.')
    candidates.sort(key=lambda c: c[0], reverse=True)
    _, name, index, row = candidates[0]
    headers = [str(v).strip() or f'Column {i + 1}' for i, v in enumerate(row)]
    if len(set(headers)) != len(headers):
        raise ValueError('Duplicate column headers found. Give each column a unique name.')
    detected = {}
    ambiguous = []
    for field, aliases in ALIASES.items():
        matches = [h for h in headers if key(h) in aliases]
        detected[field] = matches[0] if len(matches) == 1 else ''
        if len(matches) > 1:
            ambiguous.append(field)
    body = frames[name].iloc[index + 1:].copy()
    body.columns = headers
    body = body.loc[~body.apply(lambda r: all(clean(v) == '' for v in r), axis=1)]
    if len(body) > LIMIT:
        raise ValueError(f'Maximum {LIMIT:,} student rows per file.')
    return {'sheet': name, 'sheets': list(frames), 'header_row': int(index) + 1, 'columns': headers,
            'mapping': detected, 'ambiguous': ambiguous, 'rows': len(body),
            'preview': [[str(v) for v in r] for r in body.head(4).values.tolist()]}, body


def records(data, filename, config):
    info, frame = inspect(data, filename, config.get('sheet'))
    mapping = config.get('mapping') or info['mapping']
    mode = config.get('mode', 'auto')
    if mode == 'auto':
        if not mapping.get('status'):
            raise ValueError('No attendance column detected. Choose “Absentees only” if this is a complete absentee list.')
        mode = 'register'
    if mode not in {'register', 'absentees'}:
        raise ValueError('Choose a valid file type.')
    if not mapping.get('roll') or (mode == 'register' and not mapping.get('status')):
        raise ValueError('Select the roll-number column and, for a register, the attendance column.')
    for column in mapping.values():
        if column and column not in frame.columns:
            raise ValueError(f'Column “{column}” is no longer available. Inspect the file again.')
    result = {}
    duplicates = 0
    inferred = 0
    for index, row in frame.iterrows():
        roll = clean(row[mapping['roll']])
        if not roll:
            raise ValueError(f'Row {index + 1}: missing roll number. Remove this row or supply an ID.')
        if len(roll) > 100:
            raise ValueError(f'Row {index + 1}: roll number is too long.')
        status = 'Absent'
        if mode == 'register':
            value = clean(row[mapping['status']]).lower()
            if value not in PRESENT | ABSENT:
                raise ValueError(f'Row {index + 1} ({roll}): unrecognized attendance “{value}”. Use P/A, Present/Absent, Yes/No or 1/0.')
            status = 'Present' if value in PRESENT else 'Absent'
            if key(mapping['status']) == 'absent':
                # Boolean "Absent" columns mean 1/Yes = absent, not present.
                if value in {'yes', 'y', '1', 'true', 'no', 'n', '0', 'false'}:
                    status = 'Absent' if value in PRESENT else 'Present'
        batch = clean(row[mapping['batch']]) if mapping.get('batch') else ''
        branch = clean(row[mapping['branch']]) if mapping.get('branch') else ''
        batch = batch or clean(config.get('batch', ''))
        branch = branch or clean(config.get('branch', ''))
        if len(batch) > 200 or len(branch) > 200:
            raise ValueError(f'Row {index + 1}: batch or branch exceeds 200 characters.')
        if config.get('infer_ids'):
            match = re.fullmatch(r'(\d{2})([A-Z]{2,5})(\d{3,})', roll)
            if match and (not batch or not branch):
                batch, branch = batch or f'BT{match[1]}', branch or match[2]
                inferred += 1
        record = {'batch': batch, 'branch': branch, 'status': status}
        if roll in result:
            if result[roll] != record:
                raise ValueError(f'Conflicting duplicate for {roll}. Fix its attendance, batch or branch before comparing.')
            duplicates += 1
        result[roll] = record
    return result, {**info, 'mode': mode, 'unique': len(result), 'duplicates': duplicates, 'inferred': inferred}


def reconcile(cdc, classroom, cdc_mode, class_mode):
    ids = set(cdc) | set(classroom)
    if cdc_mode == 'register' and set(classroom) - set(cdc):
        raise ValueError('The CDC register is missing students found in Class. Complete the register or use the correct absentee-list mode.')
    if class_mode == 'register' and set(cdc) - set(classroom):
        raise ValueError('The Class register is missing students found in CDC. Complete the register or use the correct absentee-list mode.')
    rows = []
    for roll in sorted(ids):
        left, right = cdc.get(roll, {}), classroom.get(roll, {})
        for field in ('batch', 'branch'):
            if left.get(field) and right.get(field) and left[field] != right[field]:
                raise ValueError(f'{roll}: conflicting {field} values across the two files.')
        a, b = left.get('status', 'Present'), right.get('status', 'Present')
        missing = 'Both' if a == b == 'Absent' else 'CDC' if a == 'Absent' else 'Class' if b == 'Absent' else 'None'
        rows.append(dict(zip(COLUMNS, [left.get('batch') or right.get('batch') or 'Unspecified', left.get('branch') or right.get('branch') or 'Unspecified', roll, a, b, missing])))
    return rows


def summary(rows):
    return {'total': len(rows), **{status: sum(r['Missing Attendance'] == status for r in rows) for status in ('None', 'CDC', 'Class', 'Both')}}


def excel(rows, notes):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Attendance'
    sheet.append(COLUMNS)
    colors = {'None': 'DCFCE7', 'CDC': 'FEF3C7', 'Class': 'DBEAFE', 'Both': 'FEE2E2'}
    for row in rows:
        sheet.append([row[c] for c in COLUMNS])
        for cell in sheet[sheet.max_row]:
            # Force untrusted values to literal text, including formula-like IDs.
            cell.data_type = 's'
            cell.fill = PatternFill('solid', fgColor=colors[row['Missing Attendance']])
            cell.alignment = Alignment(vertical='center')
        sheet.row_dimensions[sheet.max_row].height = 23
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='133E38')
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip('ABCDEF', [19, 20, 26, 23, 23, 25]):
        sheet.column_dimensions[column].width = width
    if rows:
        table = Table(displayName='AttendanceTable', ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False)
        sheet.add_table(table)
    totals = workbook.create_sheet('Summary')
    totals.append(['Category', 'Count'])
    for name, count in summary(rows).items():
        totals.append([name, count])
    totals.append(['Scope', 'Exported rows only'])
    totals.append(['Meaning', 'Missing Attendance names the session(s) marked absent.'])
    totals.append(['Caution', 'Attendance differences do not establish intent or misconduct.'])
    for note in notes:
        totals.append(['Note', note])
        totals.cell(totals.max_row, 2).data_type = 's'
    totals.column_dimensions['A'].width = 24
    totals.column_dimensions['B'].width = 110
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
