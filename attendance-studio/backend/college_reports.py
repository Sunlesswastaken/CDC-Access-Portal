"""College summary reports and the auditable student/hour ledger."""
import io
from copy import copy
from datetime import date
from itertools import groupby
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph

from .college import ROOT, ROMAN

HEADERS = ['DEPARTMENT', 'HOUR', 'SECTION', 'FACULTY ID', 'FACULTY NAME', 'ABSENTEES AS PER CDC CLASSES', 'ABSENTEES POSTED BY DEPARTMENT', 'DISCREPANCY']
KEYS = ['department', 'hour', 'section', 'faculty_id', 'faculty_name', 'cdc_absent', 'department_absent', 'discrepancy']
DETAILS = ['Batch Code', 'Branch', 'Roll Number', 'Student Name', 'Section', 'Hour', 'CDC Attendance', 'Class Attendance', 'Discrepancy']
RED = 'B21F2D'


def title(report):
    return f'B.TECH {ROMAN[report["semester"]]} SEMESTER CDC AND DEPARTMENT ATTENDANCE DISCREPANCY'


def display_date(report):
    return date.fromisoformat(report['date']).strftime('%d %B, %Y')


def groups(report):
    order = {'CSE':0, 'EEE':1, 'IT':2, 'AIML':3, 'DS':4, 'ECE':5}
    rows = sorted(report['rows'], key=lambda r: (order.get(r['department'],99), r['department'], r['hour'], r['section']))
    return [list(items) for _, items in groupby(rows, key=lambda r: (r['department'], r['hour']))]


def safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


def excel_report(report):
    book = Workbook()
    sheet = book.active
    sheet.title = 'College Summary'
    logo = Image(str(ROOT / 'assets/college-header.png'))
    logo.width, logo.height = 720, 98
    sheet.add_image(logo, 'A1')
    for n in range(1, 6):
        sheet.row_dimensions[n].height = 18
    for row, value in [(6, 'CAREER DEVELOPMENT CENTER'), (7, title(report))]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = sheet.cell(row, 1, value)
        cell.font = Font(size=12 if row == 6 else 10, bold=True, color=RED if row == 6 else '111111')
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        sheet.row_dimensions[row].height = 30
    sheet.merge_cells('A8:H8'); sheet['A8'] = report['session'] + ' SESSION'
    sheet['A8'].alignment = Alignment(horizontal='center')
    sheet.merge_cells('E9:H9'); sheet['E9'] = 'Date: ' + display_date(report)
    sheet['E9'].alignment = Alignment(horizontal='right')
    if report.get('partial'):
        sheet.merge_cells('A9:D9')
        sheet['A9'] = f'Partial report: {len(report["excluded"])} IDs excluded (see Excluded Records)'
        sheet['A9'].font = Font(color=RED, size=9)
    for col, heading in enumerate(HEADERS, 1):
        cell = sheet.cell(10, col, heading)
        cell.font = Font(bold=True, size=9)
    sheet.row_dimensions[10].height = 50
    thin = Side(style='thin', color='555555')
    n = 11
    for group in groups(report):
        first = n
        for row in group:
            for col, key in enumerate(KEYS, 1):
                sheet.cell(n, col, safe(row[key]))
            sheet.cell(n, 8).font = Font(color=RED, bold=True)
            sheet.row_dimensions[n].height = 35
            n += 1
        for col in (1, 2):
            if n - first > 1:
                sheet.merge_cells(start_row=first, start_column=col, end_row=n-1, end_column=col)
        sheet.merge_cells(start_row=n, start_column=1, end_row=n, end_column=5)
        sheet.cell(n, 1, 'TOTAL')
        for col, key in [(6, 'cdc_absent'), (7, 'department_absent'), (8, 'discrepancy')]:
            sheet.cell(n, col, sum(row[key] for row in group))
        for cell in sheet[n]:
            cell.font = Font(color=RED, bold=True)
        n += 1
    for row in sheet.iter_rows(min_row=10, max_row=n-1, max_col=8):
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.merge_cells(start_row=n+1, start_column=6, end_row=n+1, end_column=8)
    sheet.cell(n+1, 6, 'DEAN – CDC').font = Font(color=RED, bold=True)
    for col, width in zip('ABCDEFGH', [15, 7, 10, 15, 30, 19, 20, 16]):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'F11'
    sheet.print_title_rows = '1:10'
    sheet.print_area = f'A1:H{n+2}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.right.text = 'DEAN – CDC'
    sheet.oddFooter.right.color = RED
    sheet.oddFooter.left.text = '&P / &N'
    ledger = book.create_sheet('Student Attendance')
    ledger.append(DETAILS)
    fills = {'None': 'E5F1E4', 'CDC': 'FFF1CB', 'Class': 'FDE6D2', 'Both': 'FADDE2'}
    for record in report['details']:
        ledger.append([safe(record.get(key, '')) for key in DETAILS])
        for cell in ledger[ledger.max_row]:
            cell.fill = PatternFill('solid', fgColor=fills[record['Missing Attendance']])
    ledger.freeze_panes = 'A2'; ledger.auto_filter.ref = ledger.dimensions
    for col, width in zip('ABCDEFGHI', [16, 12, 20, 38, 12, 10, 20, 20, 24]):
        ledger.column_dimensions[col].width = width
    for cell in ledger[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='235847')
    if report.get('excluded'):
        omissions = book.create_sheet('Excluded Records')
        fields = ['Roll Number', 'Student Name', 'Source', 'Reason']
        omissions.append(fields)
        for record in report['excluded']:
            omissions.append([safe(record[key]) for key in fields])
        omissions.freeze_panes = 'A2'
        omissions.auto_filter.ref = omissions.dimensions
        for col, width in zip('ABCD', [22, 38, 20, 100]):
            omissions.column_dimensions[col].width = width
        for cell in omissions[1]:
            cell.font = Font(bold=True, color=RED)
        for row in omissions.iter_rows(min_row=2):
            omissions.row_dimensions[row[0].row].height = 44
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    notes = book.create_sheet('Report Notes')
    for line in ['Date: ' + report['date'], 'Session: ' + report['session'], 'Discrepancy counts CDC-present students whose roll numbers appear in the Class absentee list. Class absentee counts include only the CDC cohort. Students absent in CDC are out of scope.', 'Student Attendance contains one row per CDC-present student per uploaded class hour.', 'Sources: ' + ', '.join(report['sources']), *report['warnings']]:
        notes.append([safe(line)])
    notes.column_dimensions['A'].width = 120
    for row in notes:
        row[0].alignment = Alignment(wrap_text=True); notes.row_dimensions[row[0].row].height = 40
    for tab in book:
        for row in tab:
            for cell in row:
                font = copy(cell.font)
                font.name = 'Aptos'
                if tab == sheet and 10 <= cell.row < n:
                    font.sz = 9; font.bold = True
                cell.font = font
    sheet['A6'].font = Font(name='Aptos', size=16, bold=True, color=RED)
    sheet['A7'].font = Font(name='Aptos', size=12, bold=True)
    sheet['A8'].font = Font(name='Aptos', size=12, bold=True)
    output = io.BytesIO(); book.save(output)
    return output.getvalue()


def pdf_report(report):
    if 'Aptos' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('Aptos', str(ROOT / 'assets/fonts/Aptos.ttf')))
        pdfmetrics.registerFont(TTFont('Aptos-Bold', str(ROOT / 'assets/fonts/Aptos-Bold.ttf')))
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=54, rightMargin=54, topMargin=210, bottomMargin=52, title=title(report))
    normal = ParagraphStyle('cell', fontName='Aptos-Bold', fontSize=9, leading=11, alignment=1)
    bold = ParagraphStyle('heading', parent=normal)
    name = ParagraphStyle('name', parent=normal, alignment=0)
    def paragraph(value, style=normal):
        return Paragraph(escape(str(value)), style)
    data = [[paragraph(h, bold) for h in HEADERS]]
    commands = [('GRID', (0, 0), (-1, -1), .5, colors.black), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (0, 0), (-1, -1), 'CENTER'), ('FONTSIZE', (0, 0), (-1, -1), 8), ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3), ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4), ('TEXTCOLOR', (7, 1), (7, -1), colors.HexColor('#' + RED))]
    for group in groups(report):
        start = len(data)
        for row in group:
            data.append([paragraph(row[key], name if key == 'faculty_name' else normal) if key in ('department', 'section', 'faculty_id', 'faculty_name') else row[key] for key in KEYS])
        if len(group) > 1:
            for col in (0, 1):
                commands.append(('SPAN', (col, start), (col, len(data)-1)))
        total = len(data)
        data.append(['TOTAL', '', '', '', '', *[sum(row[k] for row in group) for k in ('cdc_absent', 'department_absent', 'discrepancy')]])
        commands.extend([('SPAN', (0, total), (4, total)), ('TEXTCOLOR', (0, total), (-1, total), colors.HexColor('#' + RED)), ('FONTNAME', (0, total), (-1, total), 'Helvetica-Bold')])
        commands.append(('NOSPLIT', (0, total-1), (-1, total)))
    table = Table(data, colWidths=[70, 35, 49, 57, 79, 57, 71, 69], repeatRows=1)
    table.setStyle(TableStyle(commands))
    table.setStyle(TableStyle([('FONTNAME', (0, 0), (-1, -1), 'Aptos-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 9)]))
    story = [table]
    def page(canvas, document):
        canvas.drawImage(str(ROOT / 'assets/college-header.png'), 60, 747, width=475, height=65, preserveAspectRatio=True, mask='auto')
        canvas.setFillColor(colors.HexColor('#' + RED)); canvas.setFont('Aptos-Bold', 16)
        canvas.drawCentredString(A4[0]/2, 729, 'CAREER DEVELOPMENT CENTER')
        canvas.setFillColor(colors.black)
        heading = Paragraph(escape(title(report)), ParagraphStyle('title', fontName='Aptos-Bold', fontSize=12, leading=14, alignment=1))
        _, height = heading.wrap(A4[0]-108, 40)
        heading.drawOn(canvas, 54, 717-height)
        canvas.setFont('Aptos-Bold', 12); canvas.drawCentredString(A4[0]/2, 671, report['session'] + ' SESSION')
        canvas.setFont('Aptos-Bold', 10); canvas.drawRightString(A4[0]-54, 653, 'Date: ' + display_date(report))
        canvas.setFillColor(colors.HexColor('#' + RED)); canvas.drawRightString(A4[0]-36, 28, 'DEAN – CDC')
        canvas.setFillColor(colors.grey); canvas.setFont('Aptos', 8); canvas.drawString(36, 28, f'Page {document.page}')
    doc.build(story, onFirstPage=page, onLaterPages=page)
    return output.getvalue()
